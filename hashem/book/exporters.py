"""Export a typeset book to PDF / EPUB / DOCX / HTML.

* **PDF** is built from the rasterised pages — guaranteed-correct Arabic.
* **EPUB / DOCX / HTML** carry *logical* text so ebook readers, Word and browsers
  apply their own shaping/bidi and the text stays selectable and reflowable.
"""

from __future__ import annotations

from pathlib import Path
from typing import Sequence

from ..utils.log import get_logger
from .typesetter import TypesetResult, render_page

log = get_logger("book.exporters")

DPI = 150


def _safe_name(title: str) -> str:
    slug = "".join(ch for ch in (title or "book") if ch.isalnum() or ch in " -_").strip()
    return slug.replace(" ", "_")[:60] or "book"


def export_pdf(result: TypesetResult, output: str | Path) -> Path:
    """Rasterise every page and pack them into a real PDF."""
    from fpdf import FPDF

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    theme = result.theme
    page_mm_w = theme.page_w / DPI * 25.4
    page_mm_h = theme.page_h / DPI * 25.4

    import io

    pdf = FPDF(unit="mm", format=(page_mm_w, page_mm_h))
    pdf.set_auto_page_break(auto=False)
    for index in range(result.page_count):
        image = render_page(result, index)
        pdf.add_page()
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)
        pdf.image(buffer, x=0, y=0, w=page_mm_w, h=page_mm_h)
    pdf.output(str(output))
    log.info("PDF exported: %s (%d pages)", output, result.page_count)
    return output


def export_epub(result: TypesetResult, output: str | Path) -> Path:
    """Build an EPUB with logical text (readers do their own bidi)."""
    from ebooklib import epub

    output = Path(output)
    doc = result.doc
    book = epub.EpubBook()
    book.set_identifier(f"hashem-{_safe_name(doc.title)}")
    book.set_title(doc.title)
    book.set_language("ar" if doc.language == "ar" else "en")
    if doc.author:
        book.add_author(doc.author)

    style = epub.EpubItem(
        uid="style", file_name="style/default.css",
        media_type="text/css",
        content="body{font-family:'Amiri','Noto Naskh Arabic',serif;line-height:1.9;direction:rtl;}"
        "h1,h2{color:#333}blockquote{border-right:3px solid #b45309;padding-right:1em;color:#b45309}"
        "pre{background:#f4f4f4;padding:1em;direction:ltr;text-align:left}",
    )
    book.add_item(style)

    chapters = []
    current_blocks = []
    current_title = doc.title
    chapter_index = 0

    def flush_chapter() -> None:
        nonlocal current_blocks, chapter_index
        if not current_blocks:
            return
        chapter_index += 1
        html = [f"<h1>{_escape(current_title)}</h1>"]
        for kind, text in current_blocks:
            if kind == "heading":
                html.append(f"<h2>{_escape(text)}</h2>")
            elif kind == "subheading":
                html.append(f"<h3>{_escape(text)}</h3>")
            elif kind == "quote":
                html.append(f"<blockquote>{_escape(text)}</blockquote>")
            elif kind == "code":
                html.append(f"<pre>{_escape(text)}</pre>")
            elif kind == "list":
                html.append(f"<p>• {_escape(text)}</p>")
            elif kind == "hr":
                html.append("<hr/>")
            else:
                html.append(f"<p>{_escape(text)}</p>")
        chapter = epub.EpubHtml(
            uid=f"chapter_{chapter_index}",
            file_name=f"chapter_{chapter_index}.xhtml",
            lang=doc.language,
            content="".join(html),
        )
        chapter.add_item(style)
        book.add_item(chapter)
        chapters.append(chapter)
        current_blocks = []

    for block in doc.blocks:
        if block.kind == "heading" and current_blocks:
            flush_chapter()
            current_title = block.text
            continue
        if block.kind in {"title",}:
            continue
        current_blocks.append((block.kind, block.text))
    flush_chapter()

    book.toc = chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *chapters]
    epub.write_epub(str(output), book)
    log.info("EPUB exported: %s (%d chapters)", output, len(chapters))
    return output


def export_docx(result: TypesetResult, output: str | Path) -> Path:
    """Build a Word document with real heading styles."""
    import docx
    from docx.shared import Pt, RGBColor

    output = Path(output)
    doc = result.doc
    theme = result.theme
    document = docx.Document()
    document.add_heading(doc.title, level=0)
    if doc.author:
        paragraph = document.add_paragraph(doc.author)
        paragraph.runs[0].italic = True

    for block in doc.blocks:
        if block.kind == "title":
            continue
        if block.kind == "heading":
            document.add_heading(block.text, level=1)
        elif block.kind == "subheading":
            document.add_heading(block.text, level=2)
        elif block.kind == "quote":
            paragraph = document.add_paragraph(block.text)
            paragraph.style = "Intense Quote" if "Intense Quote" in [s.name for s in document.styles] else paragraph.style
        elif block.kind == "code":
            paragraph = document.add_paragraph(block.text)
            paragraph.style = "No Spacing"
        elif block.kind == "list":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "hr":
            document.add_paragraph("―" * 24)
        else:
            document.add_paragraph(block.text)
    document.save(str(output))
    log.info("DOCX exported: %s", output)
    return output


def export_html(result: TypesetResult, output: str | Path) -> Path:
    """A beautiful, selectable single-file HTML edition."""
    output = Path(output)
    doc = result.doc
    theme = result.theme
    parts = []
    for block in doc.blocks:
        if block.kind == "title":
            parts.append(f"<h1 class='book-title'>{_escape(block.text)}</h1>")
        elif block.kind == "heading":
            parts.append(f"<h2>{_escape(block.text)}</h2>")
        elif block.kind == "subheading":
            parts.append(f"<h3>{_escape(block.text)}</h3>")
        elif block.kind == "quote":
            parts.append(f"<blockquote>{_escape(block.text)}</blockquote>")
        elif block.kind == "code":
            parts.append(f"<pre><code>{_escape(block.text)}</code></pre>")
        elif block.kind == "list":
            parts.append(f"<p class='list'>• {_escape(block.text)}</p>")
        elif block.kind == "hr":
            parts.append("<hr/>")
        else:
            parts.append(f"<p>{_escape(block.text)}</p>")

    rtl = "rtl" if doc.language == "ar" else "ltr"
    font_stack = "'Amiri','Noto Naskh Arabic','Cairo','Segoe UI',serif" if doc.language == "ar" else "Georgia,'Segoe UI',serif"
    html = f"""<!DOCTYPE html>
<html lang="{doc.language}" dir="{rtl}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_escape(doc.title)}</title>
<style>
  :root {{ --paper:{theme.paper}; --ink:{theme.ink}; --accent:{theme.accent}; }}
  body {{ background:#111; margin:0; }}
  main {{ max-width:44em; margin:2rem auto; background:var(--paper); color:var(--ink);
         padding:3rem 2.6rem; box-shadow:0 10px 60px rgba(0,0,0,.5); line-height:2;
         font-family:{font_stack}; font-size:1.1rem; }}
  h1.book-title {{ font-size:2.4rem; text-align:center; color:var(--accent); }}
  h2 {{ color:var(--accent); border-bottom:2px solid var(--accent); padding-bottom:.3em; margin-top:2.5em; }}
  blockquote {{ border-{ 'right' if rtl=='rtl' else 'left' }:4px solid var(--accent);
      margin:1.5em 0; padding:.6em 1.2em; color:var(--accent); font-style:italic; }}
  pre {{ background:#0f172a; color:#e2e8f0; padding:1em; direction:ltr; text-align:left;
        overflow-x:auto; border-radius:8px; }}
  .list {{ padding-{ 'right' if rtl=='rtl' else 'left' }:1em; }}
  hr {{ border:0; border-top:2px solid var(--accent); width:40%; margin:2.5em auto; }}
</style>
</head>
<body><main>
<h1 class='book-title'>{_escape(doc.title)}</h1>
{'<p style="text-align:center;color:var(--accent)">' + _escape(doc.author) + '</p>' if doc.author else ''}
{''.join(parts[1:])}
</main></body></html>"""
    output.write_text(html, encoding="utf-8")
    log.info("HTML exported: %s", output)
    return output


def _escape(text: str) -> str:
    return (
        (text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


_EXPORTERS = {
    "pdf": export_pdf,
    "epub": export_epub,
    "docx": export_docx,
    "html": export_html,
}


def book_formats() -> list[str]:
    return sorted(_EXPORTERS)


def export_book(
    result: TypesetResult,
    output_dir: str | Path,
    formats: Sequence[str] = ("pdf", "html"),
) -> dict[str, Path]:
    """Export to every requested format; returns a map of format → path."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    base = _safe_name(result.doc.title)
    outputs: dict[str, Path] = {}
    for fmt in formats:
        exporter = _EXPORTERS.get(fmt.lower())
        if not exporter:
            continue
        target = output_dir / f"{base}.{fmt}"
        try:
            outputs[fmt] = exporter(result, target)
        except Exception as exc:  # noqa: BLE001 - one failing format shouldn't stop the rest
            log.exception("Export to %s failed: %s", fmt, exc)
    return outputs


__all__ = ["export_pdf", "export_epub", "export_docx", "export_html", "export_book", "book_formats"]
